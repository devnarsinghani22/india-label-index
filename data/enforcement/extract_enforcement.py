"""Build data/enforcement/fssai_enforcement.csv from the saved primary sources in data/enforcement/sources/.

Every row is copied as published (no arithmetic except where the note says so). Where two sources give different
numbers for the same year, both rows are kept and the note says so. Run from the repo root:

    python -I data/enforcement/extract_enforcement.py

Also writes data/enforcement/population_census2011.csv (Census 2011 A-1 table; Telangana and Ladakh summed from
their 2011 districts; Dadra & Nagar Haveli and Daman & Diu summed from the two 2011 UTs).
"""
import csv
import html
import json
import re
from pathlib import Path

HERE = Path(__file__).resolve().parent
SRC = HERE / "sources"
OUT = HERE / "fssai_enforcement.csv"
POP_OUT = HERE / "population_census2011.csv"

COLS = ["fiscal_year", "state", "samples_analysed", "non_conforming", "unsafe", "substandard",
        "misbranded_labelling_defects", "civil_cases", "criminal_cases", "convictions", "penalty_rs",
        "source_url", "source_page", "note"]

URLS = {
    "ls1093": "https://sansad.in/getFile/lsapps/loksabhaquestions/annex/188/AU1093_HciY1m.pdf?source=lsapps",
    "ls985": "https://sansad.in/getFile/lsapps/loksabhaquestions/annex/188/AU985_J79qdv.pdf?source=lsapps",
    "ls361": "https://fcainfoweb.nic.in/PMS/writereaddata/2025_LS_B_361.pdf",
    "ls2076_2025": "https://sansad.in/getFile/loksabhaquestions/annex/185/AU2076_kMkELH.pdf?source=pqals",
    "pib1945166": "https://pib.gov.in/PressReleasePage.aspx?PRID=1945166",
    "pib2101739": "https://pib.gov.in/PressReleasePage.aspx?PRID=2101739",
    "pib2225750": "https://pib.gov.in/PressReleasePage.aspx?PRID=2225750",
    "ar1920": "https://fssai.gov.in/docs/annual-report/FSSAI_Annual_Report_2019_20_English_Hindi.pdf",
    "ar2021": "https://fssai.gov.in/docs/annual-report/FSSAI_Annual_Report_2020-21.pdf",
    "ar2122": "https://fssai.gov.in/docs/annual-report/FSSAI_Annual_Report_2021-22.pdf",
    "datagov902": "https://www.data.gov.in/resource/stateut-wise-details-samples-analyzed-found-non-conforming-prescribed-standards-and",
}

# canonical state names (as the state column in the CSV)
CANON = {
    "andaman & nicobar island": "Andaman & Nicobar Islands", "andaman & nicobar islands": "Andaman & Nicobar Islands",
    "andaman and nicobar islands": "Andaman & Nicobar Islands", "andaman& nicobarislands": "Andaman & Nicobar Islands",
    "andhra pradesh": "Andhra Pradesh", "andhrapradesh": "Andhra Pradesh",
    "arunachal pradesh": "Arunachal Pradesh", "arunachalpradesh": "Arunachal Pradesh",
    "assam": "Assam", "bihar": "Bihar", "chandigarh": "Chandigarh", "chhattisgarh": "Chhattisgarh",
    "dadara & nagar haveli and daman & diu": "Dadra & Nagar Haveli and Daman & Diu",
    "dadra and nagar haveli & daman & diu": "Dadra & Nagar Haveli and Daman & Diu",
    "dadra nagar haveli & daman & diu": "Dadra & Nagar Haveli and Daman & Diu",
    "dadra nagar haveli& daman &diu": "Dadra & Nagar Haveli and Daman & Diu",
    "dadra nagar haveli & daman &diu": "Dadra & Nagar Haveli and Daman & Diu",
    "dadra and nagar haveli and daman and diu": "Dadra & Nagar Haveli and Daman & Diu",
    "dadara and nagar haveli and daman and diu": "Dadra & Nagar Haveli and Daman & Diu",
    "dadra & nagar haveli and daman & diu": "Dadra & Nagar Haveli and Daman & Diu",
    "nct of delhi": "Delhi", "andaman & nicobar islands ": "Andaman & Nicobar Islands",
    "jammu & kashmir ": "Jammu & Kashmir", "india": "India",
    "delhi": "Delhi", "goa": "Goa", "gujarat": "Gujarat", "haryana": "Haryana",
    "himachal pradesh": "Himachal Pradesh", "himachalpradesh": "Himachal Pradesh",
    "jammu & kashmir": "Jammu & Kashmir", "jammu&kashmir": "Jammu & Kashmir", "jammu and kashmir": "Jammu & Kashmir",
    "jharkhand": "Jharkhand", "karnataka": "Karnataka", "kerala": "Kerala", "ladakh": "Ladakh",
    "lakshadweep": "Lakshadweep", "madhya pradesh": "Madhya Pradesh", "madhyapradesh": "Madhya Pradesh",
    "maharashtra": "Maharashtra", "manipur": "Manipur", "meghalaya": "Meghalaya", "mizoram": "Mizoram",
    "nagaland": "Nagaland", "odisha": "Odisha", "orissa": "Odisha", "puducherry": "Puducherry", "punjab": "Punjab",
    "rajasthan": "Rajasthan", "sikkim": "Sikkim", "tamil nadu": "Tamil Nadu", "telangana": "Telangana",
    "tripura": "Tripura", "uttar pradesh": "Uttar Pradesh", "uttarakhand": "Uttarakhand", "west bengal": "West Bengal",
    "total": "All India",
}


def canon(name):
    key = re.sub(r"\s+", " ", name.strip().lower())
    if key not in CANON:
        raise ValueError(f"unknown state name: {name!r}")
    return CANON[key]


def num(tok):
    tok = tok.replace(",", "").strip()
    return int(tok) if tok else ""


def tokens(text):
    return re.sub(r"\s+", " ", text).split(" ")


def rows_from_tokens(toks, ncols, serial_max=36):
    """Walk tokens; a row = serial number, then name words, then exactly ncols numbers (or the word Total + ncols numbers).
    Rows with fewer numbers than ncols (blank cells in the source) are returned with partial=True."""
    out = []
    i = 0
    n = len(toks)
    while i < n:
        t = toks[i]
        is_serial = t.isdigit() and 1 <= int(t) <= serial_max and i + 1 < n and not re.fullmatch(r"[\d,]+", toks[i + 1])
        is_total = t.lower() == "total" and i + 1 < n and re.fullmatch(r"[\d,]+", toks[i + 1])
        if not (is_serial or is_total):
            i += 1
            continue
        j = i if is_total else i + 1
        name = [] if is_total else []
        if is_total:
            name = ["Total"]
            j = i + 1
        else:
            while j < n and not re.fullmatch(r"[\d,]+", toks[j]):
                name.append(toks[j])
                j += 1
        nums = []
        while j < n and re.fullmatch(r"[\d,]+", toks[j]) and len(nums) < ncols:
            # stop if this number is the serial of the next row (next token is a word)
            if len(nums) >= 2 and toks[j].isdigit() and 1 <= int(toks[j]) <= serial_max and j + 1 < n \
                    and not re.fullmatch(r"[\d,]+", toks[j + 1]) and toks[j + 1].lower() != "total":
                break
            nums.append(num(toks[j]))
            j += 1
        if nums and name:
            out.append({"name": " ".join(name), "nums": nums, "partial": len(nums) < ncols})
        i = j if j > i else i + 1
        if is_total:
            break
    return out


def page_text(txt_path, page_no):
    t = Path(txt_path).read_text(encoding="utf-8")
    parts = re.split(r"--- page (\d+)\n", t)
    pages = {int(parts[k]): parts[k + 1] for k in range(1, len(parts), 2)}
    return pages[page_no]


def row(fy, state, **kw):
    r = {c: "" for c in COLS}
    r["fiscal_year"] = fy
    r["state"] = state
    r.update(kw)
    return r


def main():
    rows = []

    # ---- Lok Sabha UQ 1093, 24 Jul 2026: state-wise 2021-22 to 2025-26 -------------------------------------------
    pages = {5: "2025-26", 6: "2024-25", 7: "2023-24", 8: "2022-23", 9: "2021-22"}
    for pg, fy in pages.items():
        txt = page_text(SRC / "LS18_S8_UQ1093_2026-07-24.txt", pg)
        body = txt.split("Convictions", 1)[-1] if "Convictions" in txt else txt.split("Conviction", 1)[-1]
        got = rows_from_tokens(tokens(body), 4)
        assert len(got) == 37, (fy, len(got))
        civil_label = "decided with penalty" if fy in ("2025-26", "2024-25", "2023-24") else "convicted with penalty"
        for g in got:
            s, nc, civ, conv = g["nums"]
            note = f"Lok Sabha UQ 1093 answered 24 Jul 2026, Annexure for {fy}; civil_cases = civil cases {civil_label}; convictions = criminal cases, no. of convictions"
            if fy == "2025-26":
                note += "; 2025-26 data provisional (source footnote)"
            rows.append(row(fy, canon(g["name"]), samples_analysed=s, non_conforming=nc, civil_cases=civ, convictions=conv,
                            source_url=URLS["ls1093"], source_page=pg, note=note))

    # ---- Lok Sabha UQ 985, 24 Jul 2026: national 2023-24 to 2025-26 incl. penalties in crore ------------------------
    txt = page_text(SRC / "LS18_S8_UQ985_2026-07-24.txt", 5)
    for m in re.finditer(r"(20\d\d-\d\d)\*?\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)\s+([\d.]+)\s+([\d,]+)\s+([\d.]+)", txt):
        fy, insp, s, nc, civ, pen_cr, conv, crim_cr = m.groups()
        pen_rs = round((float(pen_cr) + float(crim_cr)) * 1e7)
        rows.append(row(fy, "All India", samples_analysed=num(s), non_conforming=num(nc), civil_cases=num(civ), convictions=num(conv),
                        penalty_rs=pen_rs, source_url=URLS["ls985"], source_page=5,
                        note=f"Lok Sabha UQ 985 answered 24 Jul 2026, Annexure-I; inspections {insp}; civil_cases = decided with penalty; "
                             f"penalty published as Rs crore: civil {pen_cr} + criminal {crim_cr}, converted to rupees here"
                             + ("; 2025-26 provisional" if fy == "2025-26" else "")))

    # ---- Lok Sabha Starred Q 361, 26 Mar 2025 (Consumer Affairs, data from FSSAI): national 2019-20 to 2023-24 ----
    txt = page_text(SRC / "LS_361_2025-03-26_consumer_affairs.txt", 2)
    for m in re.finditer(r"(20\d\d-\d\d) (\d+) (\d+) (\d+) (\d+) (\d+) (\d+) ([\d.]+) (\d+) ([\d.]+)", txt):
        fy, s, nc, un, sub, lab, civ, pen_cr, conv, crim_cr = m.groups()
        rows.append(row(fy, "All India", samples_analysed=int(s), non_conforming=int(nc), unsafe=int(un), substandard=int(sub),
                        misbranded_labelling_defects=int(lab), civil_cases=int(civ), convictions=int(conv),
                        penalty_rs=round((float(pen_cr) + float(crim_cr)) * 1e7), source_url=URLS["ls361"], source_page=2,
                        note=f"Lok Sabha Starred Q 361 answered 26 Mar 2025 (Ministry of Consumer Affairs, figures 'as informed by FSSAI'); "
                             f"civil_cases = civil cases, no. of convictions (i.e. decided with penalty); convictions = criminal; "
                             f"penalty published as Rs crore: civil {pen_cr} + criminal {crim_cr}, converted to rupees here; "
                             f"labelling column is 'Labelling defects/Misleading/Misc.'"))

    # ---- Lok Sabha UQ 2076, 1 Aug 2025: national 2024-25 (cases launched) ----------------------------------------
    rows.append(row("2024-25", "All India", samples_analysed=194116, non_conforming=34388, civil_cases=31407,
                    source_url=URLS["ls2076_2025"], source_page=3,
                    note="Lok Sabha UQ 2076 answered 1 Aug 2025, Annexure I: 'samples lifted/analysed' 1,94,116; civil_cases = 'No. of cases launched' 31,407 (not split civil/criminal); "
                         "DISAGREES with the later 1,70,535 samples analysed for 2024-25 in UQ 985/1093 (Jul 2026) and PIB 2225750 (Feb 2026); both kept"))
    rows.append(row("2024-25", "Maharashtra", samples_analysed=8788, non_conforming=1250, civil_cases=1147,
                    source_url=URLS["ls2076_2025"], source_page=3,
                    note="Lok Sabha UQ 2076 answered 1 Aug 2025, Annexure I (Maharashtra): samples lifted/analysed 8,788; civil_cases = cases launched; "
                         "DISAGREES with 5,403 samples analysed in UQ 1093 (Jul 2026); both kept"))

    # ---- PIB 1945166, 2 Aug 2023 (Lok Sabha written reply, Consumer Affairs): national 2018-19 to 2022-23 ----------
    s = html.unescape(html.unescape((SRC / "pib_1945166_b.html").read_text(encoding="utf-8", errors="ignore")))
    t = re.sub(r"<[^>]+>", " ", s)
    t = re.sub(r"\s+", " ", t)
    seen_fy = set()
    for m in re.finditer(r"(201[89]-\d\d|202\d-\d\d) ([\d,]+) ([\d,]+) ([\d,]+) ([\d,NA]+) ([\d,]+) ([\d,]+) ([\d,NA]+) ([\d,]+)", t):
        fy, smp, nc, civ_l, civ_d, civ_conv, crim_l, crim_d, crim_conv = m.groups()
        if fy in seen_fy:  # the PIB page carries the same table more than once (print/share copies)
            continue
        seen_fy.add(fy)
        note = (f"PIB release 1945166 dated 2 Aug 2023 quoting a Lok Sabha written reply (MoS Consumer Affairs); civil_cases = civil cases launched "
                f"(decided {civ_d}, convictions {civ_conv}); criminal_cases = criminal cases launched (decided {crim_d}); convictions = criminal convictions")
        if fy == "2022-23":
            note += "; DISAGREES with 1,77,511 analysed / 44,626 non-conforming for 2022-23 in LS 361 (Mar 2025) and UQ 1093 (Jul 2026), an earlier provisional cut; both kept"
        rows.append(row(fy, "All India", samples_analysed=num(smp), non_conforming=num(nc), civil_cases=num(civ_l), criminal_cases=num(crim_l),
                        convictions=num(crim_conv), source_url=URLS["pib1945166"], source_page="table", note=note))

    # ---- PIB 2101739, 11 Feb 2025 (Rajya Sabha written reply): national 2020-21 to 2023-24, cases launched --------
    t = re.sub(r"\s+", " ", (SRC / "pib_2101739.txt").read_text(encoding="utf-8", errors="ignore"))
    seen_fy = set()
    for m in re.finditer(r"(202\d-\d\d) \| ([\d,]+) \| ([\d,]+) \| ([\d,]+) \| ([\d,]+) \|", t):
        fy, smp, nc, civ, crim = m.groups()
        if fy in seen_fy:
            continue
        seen_fy.add(fy)
        rows.append(row(fy, "All India", samples_analysed=num(smp), non_conforming=num(nc), civil_cases=num(civ), criminal_cases=num(crim),
                        source_url=URLS["pib2101739"], source_page="table",
                        note="PIB release 2101739 dated 11 Feb 2025 quoting a Rajya Sabha written reply (MoS Health); civil_cases and criminal_cases = cases launched"))

    # ---- PIB 2225750, 10 Feb 2026 (Rajya Sabha written reply): state-wise 2023-24 and 2024-25 samples/non-conf ----
    t = re.sub(r"\s+", " ", (SRC / "pib_2225750.txt").read_text(encoding="utf-8", errors="ignore"))
    seg = t[t.find("Annexure-I Details"):t.find("Annexure-II Details")]
    for m in re.finditer(r"\| (\d{1,2}) \| ([A-Za-z&' ]+?) \| ([\d,]+) \| ([\d,]+) \| ([\d,]+) \| ([\d,]+)(?= \|)", seg):
        _, name, s24, nc24, s25, nc25 = m.groups()
        for fy, a, b in (("2023-24", s24, nc24), ("2024-25", s25, nc25)):
            rows.append(row(fy, canon(name), samples_analysed=num(a), non_conforming=num(b), source_url=URLS["pib2225750"], source_page="Annexure-I",
                            note="PIB release 2225750 dated 10 Feb 2026 quoting a Rajya Sabha written reply (MoS Health), Annexure-I"))
    m = re.search(r"Total \| ([\d,]+) \| ([\d,]+) \| ([\d,]+) \| ([\d,]+) \|", seg)
    for fy, a, b in (("2023-24", m.group(1), m.group(2)), ("2024-25", m.group(3), m.group(4))):
        rows.append(row(fy, "All India", samples_analysed=num(a), non_conforming=num(b), source_url=URLS["pib2225750"], source_page="Annexure-I",
                        note="PIB release 2225750 dated 10 Feb 2026, Annexure-I total row"))
    m = re.search(r"2025-26 \| ([\d,]+) \| ([\d,]+)\*", t)
    rows.append(row("2025-26", "All India", samples_analysed=num(m.group(1)), non_conforming=num(m.group(2)), source_url=URLS["pib2225750"], source_page="Annexure-II",
                    note="PIB release 2225750 dated 10 Feb 2026, Annexure-II: FY 2025-26 'till date' (part year, provisional); the full-year figure is 2,23,808 / 40,023 in UQ 985 and 1093"))

    # ---- FSSAI Annual Report 2019-20, Table 12 (page 153 of the PDF): state-wise with unsafe/substandard/labelling
    txt = page_text(SRC / "FSSAI_Annual_Report_2019-20_p153.txt", 153)
    body = txt.split("Raised (Rs.)", 2)[-1]
    got = rows_from_tokens(tokens(body), 11)
    assert len(got) == 37, len(got)
    for g in got:
        state = canon(g["name"])
        base = dict(source_url=URLS["ar1920"], source_page=153)
        if not g["partial"]:
            s, nc, un, sub, lab, cl, cd, cpen, kl, kd, kpen = g["nums"]
            rows.append(row("2019-20", state, samples_analysed=s, non_conforming=nc, unsafe=un, substandard=sub, misbranded_labelling_defects=lab,
                            civil_cases=cl, criminal_cases=kl, convictions=kd, penalty_rs=cpen + kpen, **base,
                            note=f"FSSAI Annual Report 2019-20 Table 12; civil_cases = civil cases launched (decided {cd}); criminal_cases = criminal cases launched; "
                                 f"convictions = criminal cases decided; penalty_rs = civil {cpen} + criminal {kpen} (Rs)"
                                 + ("; TOTAL non-conforming 29,192 here vs 29,589 in LS 361 / PIB 1945166 for 2019-20; both kept" if state == "All India" else "")))
        else:
            s, nc = g["nums"][0], (g["nums"][1] if len(g["nums"]) > 1 else "")
            rows.append(row("2019-20", state, samples_analysed=s, non_conforming=nc, **base,
                            note="FSSAI Annual Report 2019-20 Table 12; row has blank cells in the source so only samples analysed and non-conforming are kept"))

    # ---- FSSAI Annual Report 2020-21, Table 14 (pages 59-60): state-wise with unsafe/substandard/labelling ---------
    t = Path(SRC / "FSSAI_Annual_Report_2020-21_p59-60.txt").read_text(encoding="utf-8")
    body = ""
    for pg in (59, 60):
        p = page_text(SRC / "FSSAI_Annual_Report_2020-21_p59-60.txt", pg)
        body += " " + p.split("Raised \n(Rs.)", 2)[-1]
    got = rows_from_tokens(tokens(body), 11)
    assert len(got) == 37, len(got)
    for g in got:
        state = canon(g["name"])
        base = dict(source_url=URLS["ar2021"], source_page="59-60")
        if not g["partial"]:
            s, nc, un, sub, lab, cl, cd, cpen, kl, kd, kpen = g["nums"]
            rows.append(row("2020-21", state, samples_analysed=s, non_conforming=nc, unsafe=un, substandard=sub, misbranded_labelling_defects=lab,
                            civil_cases=cl, criminal_cases=kl, convictions=kd, penalty_rs=cpen + kpen, **base,
                            note=f"FSSAI Annual Report 2020-21 Table 14; civil_cases = civil cases launched (decided {cd}); criminal_cases = criminal cases launched; "
                                 f"convictions = criminal cases decided; penalty_rs = civil {cpen} + criminal {kpen} (Rs)"))
        else:
            s, nc, un, sub, lab = (g["nums"] + [""] * 5)[:5]
            rows.append(row("2020-21", state, samples_analysed=s, non_conforming=nc, unsafe=un, substandard=sub, misbranded_labelling_defects=lab, **base,
                            note="FSSAI Annual Report 2020-21 Table 14; row has blank cells in the source so only the first five columns are kept"))

    # ---- FSSAI Annual Report 2021-22, Table 14 (pages 245-246): 18 numeric columns -------------------------------
    body = ""
    for pg in (245, 246):
        p = page_text(SRC / "FSSAI_Annual_Report_2021-22_p245-246.txt", pg)
        if pg == 245:
            p = p[p.find("Annual Public Laboratory"):]  # drop the administrative-setup table printed above Table 14
        body += " " + p
    got = rows_from_tokens(tokens(body), 18)
    assert len(got) == 37, len(got)
    for g in got:
        state = canon(g["name"])
        if g["partial"]:
            raise ValueError(f"2021-22 partial row {g}")
        lifted, s, awaited, nc_prev, nc_cur, un, sub, lab, cl, cd, cpen, pend_ao, pend_tr, kl, kconv, kacq, kpen, kpend = g["nums"]
        rows.append(row("2021-22", state, samples_analysed=s, non_conforming=nc_prev + nc_cur, unsafe=un, substandard=sub, misbranded_labelling_defects=lab,
                        civil_cases=cl, criminal_cases=kl, convictions=kconv, penalty_rs=cpen + kpen, source_url=URLS["ar2122"], source_page="245-246",
                        note=f"FSSAI Annual Report 2021-22 Table 14 (Annual Public Laboratory Testing Report); non_conforming = {nc_prev} (samples lifted previous year, "
                             f"report this year) + {nc_cur} (lifted and analysed this year), summed here as the report's own total does; civil_cases = civil cases launched "
                             f"(decided/convicted {cd}); criminal_cases = criminal cases launched; convictions = convicted with fine and imprisonment (acquitted {kacq}); "
                             f"penalty_rs = civil {cpen} + criminal {kpen} (Rs); samples lifted {lifted}, report awaited {awaited}"))

    # ---- data.gov.in copy of Rajya Sabha Session 265 UQ 902: state-wise 2020-21 to 2023-24 samples/non-conf/penalty
    j = json.loads((SRC / "datagov_RS265_UQ902_statewise_2020-21_to_2023-24.json").read_text(encoding="utf-8"))
    for rec in j["records"]:
        for fy in ("2020_21", "2021_22", "2022_23", "2023_24"):
            s = rec[f"_{fy}___number_of_samples_analysed"]
            nc = rec[f"_{fy}___total_number_of_samples_found_non__conforming"]
            pen = rec[f"_{fy}___total_penalties_raised__in_rs__"]
            rows.append(row(fy.replace("_", "-"), canon(rec["state_ut"]), samples_analysed=int(s), non_conforming=int(nc), penalty_rs=int(pen),
                            source_url=URLS["datagov902"], source_page="data.gov.in resource 2e34b7b8-39ff-4ada-a581-43dcfe58b97f",
                            note="data.gov.in table of Rajya Sabha Session 265 Unstarred Q 902 (published 3 Nov 2024, 2023-24 provisional); penalty_rs = total penalties raised (Rs)"))

    with OUT.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=COLS)
        w.writeheader()
        w.writerows(rows)
    print(f"wrote {OUT} rows={len(rows)}")

    # ---- population: Census 2011 A-1 ---------------------------------------------------------------------------
    import openpyxl
    wb = openpyxl.load_workbook(SRC / "census2011_A-1_population.xlsx", read_only=True)
    ws = wb[wb.sheetnames[0]]
    states, districts = {}, {}
    for r in ws.iter_rows(values_only=True):
        if not r or r[5] != "Total":
            continue
        level, name, persons = r[3], str(r[4]), r[10]
        clean = re.sub(r"[@&$*#]+$", "", name).strip()
        if level in ("STATE", "INDIA"):
            states[clean] = int(persons)
        elif level == "DISTRICT":
            districts[(r[0], clean)] = int(persons)
    TELANGANA_DISTRICTS = ["Adilabad", "Nizamabad", "Karimnagar", "Medak", "Hyderabad", "Rangareddy", "Mahbubnagar", "Nalgonda", "Warangal", "Khammam"]
    tel = sum(v for (sc, n), v in districts.items() if sc == "28" and n in TELANGANA_DISTRICTS)
    assert len([1 for (sc, n) in districts if sc == "28" and n in TELANGANA_DISTRICTS]) == 10, "Telangana districts"
    lad = sum(v for (sc, n), v in districts.items() if sc == "01" and n in ("Leh(Ladakh)", "Leh (Ladakh)", "Kargil"))
    assert lad > 0
    pop = {}
    for name, v in states.items():
        if name == "INDIA":
            pop["All India"] = v
        elif name == "DAMAN & DIU":
            continue  # merged below
        elif name == "DADRA & NAGAR HAVELI":
            pop["Dadra & Nagar Haveli and Daman & Diu"] = v + states["DAMAN & DIU"]
        else:
            pop[canon(name)] = v
    # split Telangana and Ladakh out of their 2011 parents
    pop["Telangana"] = tel
    pop["Andhra Pradesh"] = pop["Andhra Pradesh"] - tel
    pop["Ladakh"] = lad
    pop["Jammu & Kashmir"] = pop["Jammu & Kashmir"] - lad
    assert pop["All India"] == 1210854977
    with POP_OUT.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["state", "population_2011", "note"])
        for k in sorted(pop):
            note = "Census 2011 A-1 (Office of the Registrar General, https://censusindia.gov.in/nada/index.php/catalog/42526)"
            if k == "Telangana":
                note += "; sum of the 10 districts of undivided Andhra Pradesh that became Telangana in 2014"
            if k == "Andhra Pradesh":
                note += "; undivided Andhra Pradesh minus the 10 Telangana districts"
            if k == "Ladakh":
                note += "; Leh (Ladakh) + Kargil districts"
            if k == "Jammu & Kashmir":
                note += "; 2011 state minus Leh and Kargil districts"
            if k == "Dadra & Nagar Haveli and Daman & Diu":
                note += "; two 2011 UTs summed (merged 2020)"
            w.writerow([k, pop[k], note])
    print(f"wrote {POP_OUT} states={len(pop)}; All India {pop['All India']:,}; Telangana {tel:,}; Ladakh {lad:,}")


if __name__ == "__main__":
    main()
